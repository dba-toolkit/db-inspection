from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from plugins.oracle.report_adapter import REPORT_CONTRACT, adapt_report_model
from plugins.oracle.word_report import ORACLE_WORD_PROFILE, OracleWordReportGenerator


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tests" / "baselines" / "oracle" / "current" / "report_model.json"
LOGO = ROOT / "assets" / "logo.png"


def load_source() -> dict:
    return json.loads(SOURCE.read_text(encoding="utf-8"))


class OraclePluginTests(unittest.TestCase):
    def test_adapter_preserves_business_facts(self) -> None:
        source = load_source()
        result = adapt_report_model(source)
        self.assertEqual(result["generator_contract"], REPORT_CONTRACT)
        self.assertEqual(result["health_assessment"]["score"], source["health_assessment"]["score"])
        # 第⑦步 severity 二级分级后 LIBRARY_CACHE 升为 critical，健康分由 43 变为 26。
        self.assertEqual(result["health_assessment"]["score"], 26)
        self.assertEqual(len(result["risk_register"]), len(source["risk_register"]))
        self.assertEqual(len(result["risk_register"]), 10)
        self.assertEqual(len(result["inspection_sections"]), 9)
        self.assertEqual(sum(len(section["items"]) for section in result["inspection_sections"]), 52)
        self.assertEqual(len(result["appendix"]["rule_evaluations"]), 34)
        self.assertTrue(all("source_severity" in finding for finding in result["risk_register"]))

    def test_adapter_assigns_stable_unique_item_ids(self) -> None:
        result = adapt_report_model(load_source())
        item_ids = [
            item["item_id"]
            for section in result["inspection_sections"]
            for item in section["items"]
        ]
        self.assertEqual(len(item_ids), 52)
        self.assertEqual(len(set(item_ids)), 52)
        self.assertTrue(all(isinstance(value, str) and value for value in item_ids))
        self.assertEqual(item_ids[0], "system.host")
        self.assertIn("ORA.ARCHIVE.MODE", item_ids)

    def test_missing_evidence_is_not_normalized_to_passed(self) -> None:
        result = adapt_report_model(load_source())
        statuses = [
            item["analysis"]["status"]
            for section in result["inspection_sections"]
            for item in section["items"]
        ]
        # 20 fact-only items: 15 with rows become normal, 5 without rows stay not_evaluated.
        self.assertEqual(statuses.count("normal"), 35)
        self.assertEqual(statuses.count("not_evaluated"), 8)
        # LIBRARY_CACHE 在第⑦步升为 critical，对应巡检项由 attention 变 risk（3→4 / 6→5）。
        self.assertEqual(statuses.count("risk"), 4)
        self.assertEqual(statuses.count("attention"), 5)
        self.assertEqual(len(result["appendix"]["rule_evaluations"]), 34)

    def test_adapter_is_idempotent(self) -> None:
        source = load_source()
        self.assertEqual(adapt_report_model(adapt_report_model(source)), adapt_report_model(source))

    def test_word_profile_covers_all_sections(self) -> None:
        section_ids = {
            section["section_id"]
            for section in adapt_report_model(load_source())["inspection_sections"]
        }
        self.assertEqual(section_ids, set(ORACLE_WORD_PROFILE.section_order))
        self.assertEqual(len(ORACLE_WORD_PROFILE.section_order), 9)

    def test_adapter_normalizes_charts(self) -> None:
        result = adapt_report_model(load_source())
        charts = result["oracle_analysis"]["charts"]
        # 系统四张来自公共 OS 层，Oracle 专属四张留在插件：共 8 张。
        self.assertEqual([chart["chart_id"] for chart in charts], [
            "SYSTEM_CPU", "SYSTEM_MEMORY", "SYSTEM_DISK", "SYSTEM_NETWORK_REALTIME",
            "oracle_physical_io", "oracle_logical_vs_physical",
            "oracle_redo_rate", "oracle_parse_ratio",
        ])
        self.assertTrue(all(chart["status"] == "generated" for chart in charts))
        self.assertTrue(all(chart["file"].startswith("charts/") for chart in charts))
        # 图题必须命中登记表，否则会静默退化成 chart_id。
        self.assertTrue(all(chart["caption"] != chart["chart_id"] for chart in charts))
        self.assertEqual(
            {chart["chart_id"]: chart["source_scope"] for chart in charts},
            {
                "SYSTEM_CPU": "historical",
                "SYSTEM_MEMORY": "historical",
                "SYSTEM_DISK": "historical",
                "SYSTEM_NETWORK_REALTIME": "realtime_snapshot",
                "oracle_physical_io": "realtime_snapshot",
                "oracle_logical_vs_physical": "realtime_snapshot",
                "oracle_redo_rate": "realtime_snapshot",
                "oracle_parse_ratio": "realtime_snapshot",
            },
        )

    def test_professional_word_renders_sections_and_risk(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model_path = root / "report_model.json"
            model_path.write_text(
                json.dumps(adapt_report_model(load_source()), ensure_ascii=False),
                encoding="utf-8",
            )
            (root / "charts").mkdir()
            charts_source = ROOT / "tests" / "baselines" / "oracle" / "current" / "charts"
            for chart in charts_source.iterdir():
                (root / "charts" / chart.name).write_bytes(chart.read_bytes())
            output = root / "oracle.docx"
            OracleWordReportGenerator(model_path, output, logo_path=LOGO, layout="professional").build()
            generated = Document(output)
            paragraphs = [paragraph.text for paragraph in generated.paragraphs]
            text_content = "\n".join(paragraphs)
            table_text = "\n".join(
                cell.text
                for table in generated.tables
                for row in table.rows
                for cell in row.cells
            )
            for heading in ("风险登记册", "分级整改与闭环计划", "综合结论与管理建议", "附录"):
                self.assertIn(heading, text_content)
            for finding in adapt_report_model(load_source())["risk_register"]:
                self.assertIn(finding["finding_id"], table_text)
                self.assertIn(finding["title"], table_text)
            self.assertIn("26 / 100", table_text)
            self.assertIn("NOARCHIVELOG", table_text)
            # 8 张技术图 + 封面 1 个 logo；逐张都从 charts/ 取到才会被内嵌。
            self.assertEqual(len(generated.inline_shapes), 9)


if __name__ == "__main__":
    unittest.main()
