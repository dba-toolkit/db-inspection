from __future__ import annotations

import hashlib
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from docx import Document
from inspection_core.package_io import read_json
from inspection_core.word_engine import WordReportEngine
from plugins.mysql.word_report import MYSQL_WORD_PROFILE, MySQLWordReportGenerator


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_MODEL = PROJECT_ROOT / "tests" / "baselines" / "mysql" / "current" / "report_model.json"
LOGO = PROJECT_ROOT / "logo.png"
PRE_SPLIT_NORMALIZED_DOCUMENT_XML = "F7B0D2457FB0F1BDB8BAF4E36263EEC17AC87884BF145871A7AAFF650543C915"


def normalized_document_xml_hash(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    for node in root.iter():
        node.attrib.pop("descr", None)
        node.attrib.pop("title", None)
    return hashlib.sha256(ET.tostring(root, encoding="utf-8")).hexdigest().upper()


class WordEngineTests(unittest.TestCase):
    def test_mysql_entry_is_a_thin_shared_engine_adapter(self) -> None:
        self.assertTrue(issubclass(MySQLWordReportGenerator, WordReportEngine))
        self.assertEqual(MySQLWordReportGenerator.__module__, "plugins.mysql.word_report")

    def test_shared_engine_has_no_database_plugin_dependency(self) -> None:
        source = (PROJECT_ROOT / "inspection_core" / "word_engine.py").read_text(encoding="utf-8")
        self.assertNotIn("plugins.mysql", source)
        self.assertNotIn("mysql_inspection_report_model", source)

    def test_mysql_profile_owns_contract_sections_and_charts(self) -> None:
        self.assertEqual(MYSQL_WORD_PROFILE.contracts, ("mysql_inspection_report_model",))
        self.assertEqual(len(MYSQL_WORD_PROFILE.section_order), 9)
        self.assertEqual(set(MYSQL_WORD_PROFILE.section_charts), {"system_performance", "mysql_runtime"})
        chart_ids = {
            definition.chart_id
            for section in MYSQL_WORD_PROFILE.section_charts.values()
            for definition in section.charts
        }
        self.assertEqual(len(chart_ids), 6)

    def test_generated_document_matches_pre_split_content_and_has_image_alt_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "report.docx"
            MySQLWordReportGenerator(
                REPORT_MODEL,
                output,
                author="王劲松",
                reviewer="邓秋爽",
                logo_path=LOGO,
                company=MYSQL_WORD_PROFILE.default_company,
                layout="legacy",
            ).build()
            generated = Document(output)
            self.assertEqual(
                normalized_document_xml_hash(output),
                PRE_SPLIT_NORMALIZED_DOCUMENT_XML,
            )
            self.assertEqual(len(generated.inline_shapes), 7)
            self.assertTrue(
                all(shape._inline.docPr.get("descr") for shape in generated.inline_shapes)
            )

    def test_professional_layout_places_facts_before_chapter_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "professional.docx"
            MySQLWordReportGenerator(
                REPORT_MODEL,
                output,
                author="王劲松",
                reviewer="邓秋爱",
                logo_path=LOGO,
                company=MYSQL_WORD_PROFILE.default_company,
                layout="professional",
            ).build()
            generated = Document(output)
            paragraphs = [paragraph.text for paragraph in generated.paragraphs]
            first_item = paragraphs.index("3.1 主机与操作系统信息")
            chapter_analysis = paragraphs.index("3.7 本章分析结论")
            self.assertLess(first_item, chapter_analysis)
            self.assertFalse(
                any(value.startswith("分析结论（") for value in paragraphs[first_item:chapter_analysis])
            )
            text_content = "\n".join(paragraphs)
            self.assertIn("12. 风险登记册", text_content)
            self.assertIn("13. 分级整改与闭环计划", text_content)
            self.assertIn("14. 综合结论与管理建议", text_content)
            self.assertIn("15. 附录", text_content)
            self.assertIn("未登记", text_content)
            table_text = "\n".join(
                cell.text
                for table in generated.tables
                for row in table.rows
                for cell in row.cells
            )
            model = read_json(REPORT_MODEL)
            self.assertIn(f"{model['health_assessment']['score']} / 100", table_text)
            for finding in model["risk_register"]:
                self.assertIn(finding["finding_id"], table_text)
                self.assertIn(finding["title"], table_text)
                self.assertIn(finding["recommendation"], table_text)
            with zipfile.ZipFile(output) as archive:
                footer_xml = b"\n".join(
                    archive.read(name)
                    for name in archive.namelist()
                    if name.startswith("word/footer") and name.endswith(".xml")
                )
            self.assertIn(b"NUMPAGES", footer_xml)

    def test_professional_layout_rejects_unknown_mode(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaisesRegex(ValueError, "layout"):
                MySQLWordReportGenerator(
                    REPORT_MODEL,
                    Path(temporary) / "invalid.docx",
                    layout="unknown",
                )

    def test_professional_layout_uses_black_hierarchy_and_reader_friendly_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "professional.docx"
            MySQLWordReportGenerator(
                REPORT_MODEL,
                output,
                logo_path=LOGO,
                layout="professional",
            ).build()
            generated = Document(output)
            self.assertEqual(str(generated.styles["Heading 1"].font.color.rgb), "1A1A1A")
            self.assertEqual(str(generated.styles["Heading 2"].font.color.rgb), "1A1A1A")
            self.assertEqual(str(generated.styles["Heading 3"].font.color.rgb), "222222")
            self.assertNotIn("-", [paragraph.text for paragraph in generated.paragraphs])
            all_rows = [
                [cell.text for cell in row.cells]
                for table in generated.tables
                for row in table.rows
            ]
            generated_time = next(row[1] for row in all_rows if row and row[0] == "生成时间")
            collection_time = next(row[1] for row in all_rows if row and row[0] == "采集时间")
            self.assertTrue(generated_time.endswith("（本地时间）"))
            self.assertTrue(collection_time.endswith("（本地时间）"))
            self.assertNotIn("+0800", generated_time + collection_time)
            self.assertFalse(any(row[:2] == ["动作", "要求"] for row in all_rows))
            self.assertTrue(any("无记录" in cell for row in all_rows for cell in row))


if __name__ == "__main__":
    unittest.main()
