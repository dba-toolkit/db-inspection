from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from plugins.sqlserver.report_adapter import REPORT_CONTRACT, adapt_report_model
from plugins.sqlserver.word_report import SQLSERVER_WORD_PROFILE, SQLServerWordReportGenerator


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tests" / "baselines" / "sqlserver" / "current" / "report_model.json"
LOGO = ROOT / "assets" / "logo.png"


def load_source() -> dict:
    return json.loads(SOURCE.read_text(encoding="utf-8"))


def _full_snapshot() -> dict:
    """一个每章都能产出条目的最小快照。

    用途是验"章节顺序契约"：基线包缺 HA / os_metrics 等数据，走基线只能覆盖 12 章，
    漏掉的正是"未登记 -> 排到最后"的那一类。这里每个键只给一行，够触发即可。
    """
    return {
        "schema_version": "1.1",
        "target": {"server": "192.168.1.37"},
        "instance": {"server_name": "SQLDB01", "product_version": "15.0.2000.5"},
        "databases": [{"name": "Eureka_NIS", "state_desc": "ONLINE"}],
        "backups": [{"database_name": "Eureka_NIS", "recovery_model_desc": "FULL"}],
        "buffer_pool": [{"database_name": "Eureka_NIS", "cache_size_mb": 100.0, "cache_pct": 50.0}],
        "volumes": [{"volume_mount_point": "D:\\", "available_mb": 1000.0}],
        "performance_samples": [{"batch_requests_per_sec": 10.0, "user_connections": 5}],
        "wait_stats": [{"wait_type": "PAGEIOLATCH_SH", "wait_time_ms": 100.0}],
        "tempdb": {"total_mb": 1024.0, "used_mb": 128.0, "used_pct": 12.5, "files": [{}]},
        "top_queries": [{"database_name": "Eureka_NIS", "execution_count": 3}],
        "tables_without_pk": [{"database_name": "Eureka_NIS", "table_count": 1}],
        "security": {"security_checks": [{"check_name": "sa 账号状态", "result": "PASS"}],
                     "sql_logins": [{"name": "sa"}], "sysadmin_members": [{"name": "sa"}]},
        "availability_replicas": [{"availability_group": "AG1", "replica_server_name": "SQLDB02",
                                   "database_name": "Eureka_NIS", "role_desc": "SECONDARY",
                                   "availability_mode_desc": "SYNCHRONOUS_COMMIT",
                                   "synchronization_state_desc": "SYNCHRONIZED",
                                   "synchronization_health_desc": "HEALTHY"}],
        "replication": [{"record_type": "PUBLICATION", "name": "Pub1", "status_desc": "TRANSACTIONAL"}],
        "failed_jobs": [{"job_name": "BackupFull", "run_datetime": "2026-09-21T01:00:00"}],
        "configurations": [{"name": "max server memory (mb)", "value_in_use": 8192}],
    }


class SQLServerPluginTests(unittest.TestCase):
    def test_adapter_preserves_business_facts(self) -> None:
        source = load_source()
        result = adapt_report_model(source)
        self.assertEqual(result["generator_contract"], REPORT_CONTRACT)
        self.assertEqual(result["health_assessment"]["score"], source["health_assessment"]["score"])
        self.assertEqual(result["health_assessment"]["score"], 0)
        self.assertEqual(len(result["risk_register"]), len(source["risk_register"]))
        self.assertEqual(len(result["risk_register"]), 24)
        self.assertEqual(len(result["inspection_sections"]), 12)
        self.assertEqual(sum(len(section["items"]) for section in result["inspection_sections"]), 27)
        self.assertEqual(len(result["appendix"]["rule_evaluations"]), 14)
        self.assertTrue(all("source_severity" in finding for finding in result["risk_register"]))

    def test_adapter_assigns_stable_unique_item_ids(self) -> None:
        result = adapt_report_model(load_source())
        item_ids = [
            item["item_id"]
            for section in result["inspection_sections"]
            for item in section["items"]
        ]
        self.assertEqual(len(item_ids), 27)
        self.assertEqual(len(set(item_ids)), 27)
        self.assertTrue(all(isinstance(value, str) and value for value in item_ids))
        self.assertEqual(item_ids[0], "instance.01")

    def test_headers_and_rows_are_normalized_to_table_rows(self) -> None:
        result = adapt_report_model(load_source())
        first_item = result["inspection_sections"][0]["items"][0]
        rows = first_item["display"]["rows"]
        self.assertTrue(rows)
        self.assertTrue(all(isinstance(row, dict) for row in rows))
        self.assertEqual(first_item["collection"]["status"], "ok")

    def test_missing_evidence_is_not_normalized_to_passed(self) -> None:
        result = adapt_report_model(load_source())
        statuses = [
            item["analysis"]["status"]
            for section in result["inspection_sections"]
            for item in section["items"]
        ]
        self.assertEqual(statuses.count("normal"), 16)
        self.assertEqual(statuses.count("risk"), 6)
        self.assertEqual(statuses.count("attention"), 5)
        self.assertEqual(len(result["appendix"]["rule_evaluations"]), 14)

    def test_management_fields_survive_in_remediation_plan(self) -> None:
        result = adapt_report_model(load_source())
        plan = result["optimization_plan"]
        self.assertEqual({key: len(value) for key, value in plan.items()}, {"P1": 17, "P2": 6, "P3": 1})
        first = plan["P1"][0]
        self.assertIn("owner", first)
        self.assertIn("planned_window", first)
        self.assertIn("verification", first)

    def test_word_profile_covers_all_sections(self) -> None:
        """基线包产出的章节必须都在 section_order 里登记。

        这里刻意用子集断言而不是相等断言：适配器会跳过无条目的章节，而基线包恰好没有
        HA/复制数据，若按"集合相等"验，'ha' 不登记反而成了通过条件 —— 它此前正是这么
        漏掉 ha 的。双向校验在 test_word_profile_covers_every_emittable_section。
        """
        section_ids = {
            section["section_id"]
            for section in adapt_report_model(load_source())["inspection_sections"]
        }
        registered = set(SQLSERVER_WORD_PROFILE.section_order)
        self.assertTrue(
            section_ids <= registered,
            f"基线包章节未在 section_order 登记：{sorted(section_ids - registered)}",
        )
        self.assertEqual(len(SQLSERVER_WORD_PROFILE.section_order), 14)

    def test_word_profile_covers_every_emittable_section(self) -> None:
        """章节顺序契约必须按"呈现层能产出的全部章节"验，不能只验基线包。

        适配器会跳过无条目的章节，而基线包恰好没有 HA/复制数据、没有 os_metrics，
        所以 ha 章节在基线里是空的 —— 只拿适配器输出比集合，等于给"ha 没登记"
        这个缺口背书：真实带 AlwaysOn 的包一进来，未登记章节会被排到最后（顺序 99）。
        """
        from plugins.sqlserver.presentation import build_model

        snapshot = _full_snapshot()
        model = build_model(snapshot, [], [])
        emitted = {section["section_id"] for section in model["sections"]}
        registered = set(SQLSERVER_WORD_PROFILE.section_order)
        self.assertTrue(
            emitted <= registered,
            f"呈现层产出的章节未在 section_order 登记，会被排到附录前：{sorted(emitted - registered)}",
        )
        # 反向也要成立：登记了却永不产出的章节是死配置，会让人误以为报告有这一章。
        self.assertTrue(
            registered <= emitted,
            f"section_order 登记了呈现层从不产出的章节：{sorted(registered - emitted)}",
        )

    def test_professional_word_renders_sections_risk_and_charts(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model_path = root / "report_model.json"
            model_path.write_text(
                json.dumps(adapt_report_model(load_source()), ensure_ascii=False),
                encoding="utf-8",
            )
            (root / "charts").mkdir()
            charts_source = ROOT / "tests" / "baselines" / "sqlserver" / "current" / "charts"
            for chart in charts_source.iterdir():
                (root / "charts" / chart.name).write_bytes(chart.read_bytes())
            output = root / "sqlserver.docx"
            SQLServerWordReportGenerator(model_path, output, logo_path=LOGO, layout="professional").build()
            generated = Document(output)
            text_content = "\n".join(paragraph.text for paragraph in generated.paragraphs)
            table_text = "\n".join(
                cell.text
                for table in generated.tables
                for row in table.rows
                for cell in row.cells
            )
            for heading in ("风险登记册", "分级整改与闭环计划", "综合结论与管理建议", "附录"):
                self.assertIn(heading, text_content)
            for finding in adapt_report_model(load_source())["risk_register"]:
                self.assertIn(finding["finding_id"], table_text)
            self.assertIn("完整备份超期或缺失", table_text)
            self.assertGreaterEqual(len(generated.inline_shapes), 9)


if __name__ == "__main__":
    unittest.main()
