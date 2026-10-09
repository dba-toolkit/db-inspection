"""SQL Server 报告文案层（analysis_text）的守门测试。

三件事必须同时成立，报告才不会退回「表格清单」：

1. **结构对账**：``presentation.py`` 声明的每个检查项，都在
   ``analysis_text`` 的注册表里有对应文案 —— 少一条就说明有人加了检查项却没写结论，
   那一项会重新落回引擎的兜底文案（这正是本次修复要消灭的形态）。
   反过来注册表也不能有死键（改名后残留的条目）。
2. **运行时覆盖**：真包跑出来的每个检查项都必须有非空的分析结论、判断依据与口径说明，
   且不得出现旧的兜底套话。
3. **口径一致**：报告模型里的 ``scoring_policy`` 必须与 ``metrics.SEVERITY_WEIGHTS``
   同源 —— 首页那段扣分算式的全部价值就是「读者能拿它把分数核对一遍」。
"""

from __future__ import annotations

import ast
import json
import tempfile
import unittest
from pathlib import Path

from plugins.sqlserver import analysis_text
from plugins.sqlserver.metrics import SEVERITY_WEIGHTS
from plugins.sqlserver.presentation import build_model
from plugins.sqlserver.report_adapter import adapt_report_model

ROOT = Path(__file__).resolve().parents[1]
PRESENTATION = ROOT / "plugins" / "sqlserver" / "presentation.py"
RULES = ROOT / "plugins" / "sqlserver" / "inspection_rules.json"
BASELINE = ROOT / "tests" / "baselines" / "sqlserver" / "current" / "baseline.json"

FALLBACK = "不单独形成风险判断"


def _literal(dict_node: ast.Dict, key: str) -> str | None:
    for k, v in zip(dict_node.keys, dict_node.values):
        if isinstance(k, ast.Constant) and k.value == key and isinstance(v, ast.Constant):
            return str(v.value)
    return None


def _declared_items() -> set[tuple[str, str]]:
    """静态提取 presentation._build_sections 声明的全部 (section_id, title)。

    不依赖运行数据：靠读源码，所以哪怕某个分支在这次采集里没被走到，
    它声明的检查项也照样参与对账。
    """
    tree = ast.parse(PRESENTATION.read_text(encoding="utf-8"))
    fn = next(
        node for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef) and node.name == "_build_sections"
    )
    appends = [
        node for node in ast.walk(fn)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and node.func.attr == "append"
        and isinstance(node.func.value, ast.Name)
        and node.args
        and isinstance(node.args[0], ast.Dict)
    ]
    # 第一遍：建立 变量名 -> section_id（s1/s2/... 是在各自章节末尾被挂到 sections 上的）
    var_to_section: dict[str, str] = {}
    inline: set[tuple[str, str]] = set()
    for node in appends:
        if node.func.value.id != "sections":  # type: ignore[attr-defined]
            continue
        arg = node.args[0]
        section_id = _literal(arg, "section_id")
        if not section_id:
            continue
        items = next(
            (v for k, v in zip(arg.keys, arg.values)
             if isinstance(k, ast.Constant) and k.value == "items"),
            None,
        )
        if isinstance(items, ast.Name):
            var_to_section[items.id] = section_id
        elif isinstance(items, ast.List):
            for element in items.elts:
                if isinstance(element, ast.Dict):
                    title = _literal(element, "title")
                    if title:
                        inline.add((section_id, title))
    # 第二遍：逐条 sN.append({"title": ...})
    declared = set(inline)
    for node in appends:
        name = node.func.value.id  # type: ignore[attr-defined]
        if name not in var_to_section:
            continue
        title = _literal(node.args[0], "title")
        if title:
            declared.add((var_to_section[name], title))
    return declared


class DeclaredItemCoverageTests(unittest.TestCase):
    """结构对账：声明与文案注册表必须一一对应。"""

    def test_every_declared_item_has_a_spec(self) -> None:
        missing = sorted(_declared_items() - analysis_text.spec_keys())
        self.assertEqual(missing, [], f"以下检查项没有分析结论文案: {missing}")

    def test_no_orphan_specs(self) -> None:
        orphan = sorted(analysis_text.spec_keys() - _declared_items())
        self.assertEqual(orphan, [], f"以下分析结论文案没有对应检查项（改名后的残留？）: {orphan}")

    def test_declared_items_are_actually_found(self) -> None:
        # 防止 AST 提取本身失效导致前两条测试变成空对空。
        self.assertGreater(len(_declared_items()), 40)


class DecorationBehaviourTests(unittest.TestCase):
    """装修逻辑本身：成功补文案、未采集不写假结论、截断提示只在真截断时出现。"""

    def _section(self, status: str, row_count: int = 1):
        item = {
            "title": "实例基础信息",
            "source": "SERVERPROPERTY",
            "collection": {"status": status, "row_count": row_count},
            "display": {"headers": ["属性", "值"], "rows": [["a", "b"]] * row_count},
        }
        return [{"section_id": "instance", "title": "实例与数据库概览", "items": [item]}], item

    def test_missing_module_status_falls_back_to_not_evaluated(self) -> None:
        sections, item = self._section("not_evaluated", row_count=0)
        analysis_text.decorate_sections(sections, {"instance": {}})
        self.assertEqual(item["analysis"]["status"], "not_evaluated")
        self.assertIn("未取得有效结果", item["analysis"]["conclusion"])
        # 采集失败时不得残留「已取得…」这类事实陈述
        self.assertNotIn("已取得", item["analysis"]["conclusion"])

    def test_successful_collection_gets_full_analysis(self) -> None:
        sections, item = self._section("success")
        analysis_text.decorate_sections(sections, {"instance": {}, "collection": {}})
        self.assertTrue(item["analysis"]["conclusion"])
        self.assertNotIn(FALLBACK, item["analysis"]["conclusion"])
        self.assertTrue(item["display"]["note"])

    def test_truncation_note_only_when_rows_were_cut(self) -> None:
        sections, item = self._section("success", row_count=90)
        item["display"]["rows"] = [["a", "b"]] * 30
        analysis_text.decorate_sections(sections, {"instance": {}, "collection": {}})
        self.assertIn("正文展示前 30 条，共 90 条", item["display"]["note"])

        sections2, item2 = self._section("success", row_count=30)
        item2["display"]["rows"] = [["a", "b"]] * 30
        analysis_text.decorate_sections(sections2, {"instance": {}, "collection": {}})
        self.assertNotIn("正文展示前", item2["display"]["note"])

    def test_unknown_item_is_left_untouched(self) -> None:
        item = {"title": "不存在的检查项", "collection": {"status": "success"}, "display": {}}
        sections = [{"section_id": "instance", "items": [item]}]
        analysis_text.decorate_sections(sections, {})
        self.assertNotIn("analysis", item)


class ScoringPolicyTests(unittest.TestCase):
    """扣分口径必须与评分函数同源，否则首页那段话就是一段看起来精确的错话。"""

    def setUp(self) -> None:
        self.model = build_model({"collection": {"module_status": []}}, [], [])

    def test_policy_matches_score_weights(self) -> None:
        policy = self.model["health_assessment"]["scoring_policy"]
        self.assertEqual(policy["base"], 100)
        self.assertEqual(policy["minimum"], 0)
        self.assertEqual(policy["penalty_per_finding"], SEVERITY_WEIGHTS)

    def test_counts_key_present_for_shared_engine(self) -> None:
        # 引擎读的是 counts（不是 severity_counts）；缺了首页项数会全印成 0
        health = self.model["health_assessment"]
        self.assertIn("counts", health)
        self.assertEqual(health["counts"], health["severity_counts"])


class PendingConfirmationTests(unittest.TestCase):
    """待客户确认事项：只收「必须客户拍板」的规则，且模板里不能有对不上的规则号。"""

    def setUp(self) -> None:
        self.registry = json.loads(RULES.read_text(encoding="utf-8"))["rules"]

    def test_templates_reference_real_rules(self) -> None:
        unknown = sorted(set(analysis_text.CONFIRMATION_TEMPLATES) - set(self.registry))
        self.assertEqual(unknown, [], f"模板引用了规则包里不存在的规则: {unknown}")

    def test_triggered_rule_enters_the_list(self) -> None:
        findings = [{"rule_id": "BACKUP.NO_BACKUP", "evidence": "恢复模式=FULL"}]
        items = analysis_text.pending_confirmations(findings, {"server_name": "db02"})
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["topic"], analysis_text.CONFIRMATION_TEMPLATES["BACKUP.NO_BACKUP"][0])
        self.assertEqual(items[0]["evidence"], ["恢复模式=FULL"])

    def test_same_rule_declared_once(self) -> None:
        findings = [{"rule_id": "DB.CHECKDB"}, {"rule_id": "DB.CHECKDB"}, {"rule_id": "DB.RCSI"}]
        items = analysis_text.pending_confirmations(findings, {})
        self.assertEqual([i["rule_id"] for i in items], ["DB.CHECKDB", "DB.RCSI"])

    def test_unrelated_rule_does_not_enter(self) -> None:
        # 技术性整改项（容量/参数/对象治理）不进本章 —— 避免把「欢迎确认」清单空转出来
        findings = [{"rule_id": "INSTANCE.MAXDOP"}, {"rule_id": "STORAGE.FREE_SPACE"}]
        self.assertEqual(analysis_text.pending_confirmations(findings, {}), [])

    def test_no_findings_yields_empty_list_not_none(self) -> None:
        # 空列表与 None 在引擎里渲染不同：前者写「本次没有待确认事项」，后者只能说「未产出清单」
        self.assertEqual(analysis_text.pending_confirmations([], {}), [])


class AdapterPassthroughTests(unittest.TestCase):
    """适配器会重新拼装模型对象，新字段必须显式透传。

    「分析层算了、报告里没有」是这套流水线最常见的断点（适配器不认识的键会被静默丢掉），
    所以口径字段与待确认清单都要在这里钉住 —— 只测 presentation 的输出是不够的。
    """

    POLICY = {
        "base": 100,
        "penalty_per_finding": {"critical": 18, "high": 10, "medium": 4, "low": 1},
        "minimum": 0,
    }

    @staticmethod
    def _source(**extra) -> dict:
        source = {
            "environment": {"instance": {"server_name": "db02"}, "target": {"port": 1433}},
            "overview": {},
            "health_assessment": {
                "score": 0, "grade": "高风险",
                "severity_counts": {"critical": 1, "high": 4, "medium": 9, "low": 16},
                "counts": {"critical": 1, "high": 4, "medium": 9, "low": 16},
                "scoring_policy": AdapterPassthroughTests.POLICY,
                "risk_count": 30,
            },
            "collection_quality": {"score_pct": 100.0},
            "risk_register": [],
            "charts": [],
            "sections": [],
        }
        source.update(extra)
        return source

    def test_scoring_policy_survives_the_adapter(self) -> None:
        adapted = adapt_report_model(self._source())
        self.assertEqual(adapted["health_assessment"]["scoring_policy"], self.POLICY)

    def test_confirmations_survive_the_adapter(self) -> None:
        confirmations = [{"topic": "备份归属", "question": "由谁发起？", "evidence": ["恢复模式=FULL"]}]
        adapted = adapt_report_model(self._source(pending_confirmations=confirmations))
        self.assertEqual(adapted["pending_confirmations"], confirmations)

    def test_absent_confirmations_stay_none(self) -> None:
        # 旧模型没有这个键时必须保持 None：引擎对「没给字段」与「给了空列表」的措辞不同，
        # 前者只能声明「未产出清单」，后者才敢写「本次没有待确认事项」。
        adapted = adapt_report_model(self._source())
        self.assertIsNone(adapted["pending_confirmations"])


@unittest.skipUnless(BASELINE.exists(), "缺少 SQL Server 基线")
class RealPackageTextCoverageTests(unittest.TestCase):
    """真包运行时覆盖：每个检查项都得有结论、依据与口径说明。"""

    @classmethod
    def setUpClass(cls) -> None:
        baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
        package = Path(baseline["source_package"]["path"])
        if not package.exists():
            raise unittest.SkipTest(f"缺少真实采集包: {package}")
        snapshot = cls._snapshot_of(package)
        cls.model = build_model(snapshot, [], [])

    @staticmethod
    def _snapshot_of(package: Path) -> dict:
        import zipfile

        with zipfile.ZipFile(package) as archive:
            name = next(n for n in archive.namelist() if n.endswith("snapshot.json"))
            return json.loads(archive.read(name).decode("utf-8-sig"))

    def test_every_rendered_item_has_analysis_and_note(self) -> None:
        items = [item for section in self.model["sections"] for item in section["items"]]
        self.assertGreater(len(items), 20)
        for item in items:
            analysis = item.get("analysis") or {}
            with self.subTest(item=item.get("title")):
                self.assertTrue(analysis.get("conclusion"), "缺少分析结论")
                self.assertTrue(analysis.get("evidence"), "缺少判断依据")
                self.assertTrue((item.get("display") or {}).get("note"), "缺少口径说明")

    def test_no_fallback_boilerplate(self) -> None:
        items = [item for section in self.model["sections"] for item in section["items"]]
        offenders = [
            item.get("title") for item in items
            if FALLBACK in str((item.get("analysis") or {}).get("conclusion") or "")
        ]
        self.assertEqual(offenders, [], f"仍在使用兜底套话: {offenders}")

    def test_rendered_items_are_all_registered(self) -> None:
        items = [item for section in self.model["sections"] for item in section["items"]]
        for section in self.model["sections"]:
            for item in section["items"]:
                key = (str(section.get("section_id")), str(item.get("title")))
                if key not in analysis_text.spec_keys():
                    with self.subTest(item=key):
                        self.fail(f"渲染出的检查项没有登记文案: {key}")
        self.assertTrue(items)


class SingleInstanceSummaryTests(unittest.TestCase):
    """单实例报告同样要有「1.2 高风险项速览」。

    最初的筛选条件是 ``severity == "high" and finding.get("node")`` —— node 只在多
    实例输入时注入，于是单实例报告整节不渲染。这一节和节点数量无关，只和「最要紧
    的事情有没有被放在第一页」有关。本文件用的这份标准模型正是单实例、0 条带 node。
    """

    MODEL = ROOT / "tests" / "baselines" / "sqlserver" / "current" / "report_model_standard.json"

    @classmethod
    def setUpClass(cls) -> None:
        if not cls.MODEL.exists():
            raise unittest.SkipTest("缺少 SQL Server 标准报告模型")
        from docx import Document

        from plugins.sqlserver.word_report import SQLServerWordReportGenerator

        cls._tmp = tempfile.TemporaryDirectory()
        out = Path(cls._tmp.name) / "report.docx"
        SQLServerWordReportGenerator(
            cls.MODEL, out, author="测试", company="测试", logo_path=ROOT / "assets" / "logo.png",
        ).build()
        doc = Document(out)
        cls.paragraphs = [p.text.strip() for p in doc.paragraphs]
        cls.tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in doc.tables]

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def _speed_view(self) -> list[list[str]] | None:
        return next((t for t in self.tables if t and any("风险项" in c for c in t[0])), None)

    def test_speed_view_is_rendered_for_single_instance(self) -> None:
        self.assertIn("1.2 高风险项速览", self.paragraphs)

    def test_speed_view_has_no_node_column(self) -> None:
        table = self._speed_view()
        self.assertIsNotNone(table, "没有找到高风险项速览表")
        self.assertNotIn("节点", table[0])
        self.assertIn("关键事实", table[0])

    def test_speed_view_matches_homepage_high_risk_count(self) -> None:
        """速览条数必须与首页「高风险」列同口径（critical 计入高风险档）。

        不同口径就会出现「首页写 5 条、速览列 4 条」的自相矛盾。
        """
        model = json.loads(self.MODEL.read_text(encoding="utf-8"))
        expect = sum(
            1 for x in model.get("risk_register") or []
            if str(x.get("severity")) in {"critical", "high"}
        )
        self.assertGreater(expect, 0)
        table = self._speed_view()
        self.assertIsNotNone(table)
        self.assertEqual(len(table) - 1, expect)

    def test_facts_column_is_not_the_placeholder(self) -> None:
        table = self._speed_view()
        self.assertIsNotNone(table)
        facts = [row[-1] for row in table[1:]]
        self.assertTrue(facts)
        self.assertTrue(all(f and f != "未提供事实" for f in facts), facts)

    def test_scope_note_has_no_hardcoded_chapter_number(self) -> None:
        """口径说明里的章节号一旦写死就会错（风险登记册 MySQL 是 12 章、SQL Server 是 17 章）。"""
        note = next((p for p in self.paragraphs if "仅列风险登记册中的" in p), None)
        self.assertIsNotNone(note, "没有找到速览口径说明")
        self.assertNotIn("第 12 章", note)


class EmptyItemRenderingTests(unittest.TestCase):
    """空数据（success_no_data / 未采集）也不能让检查项整项消失。

    「采集成功、确认没有」与「没采到、无法判定」是两个相反的结论，在报告里不能长得
    一样，更不能一起消失 —— 与「未采集 ≠ 0」是同一条规矩。
    """

    MUST_ALWAYS_RENDER = {
        "缺失索引候选",
        "未使用索引",
        "陈旧统计信息",
        "索引碎片",
        "孤立用户",
        "可疑页面",
    }

    def _conditional_item_titles(self) -> set[str]:
        """静态找出所有「写在 if 体内的检查项 append」。"""
        tree = ast.parse(PRESENTATION.read_text(encoding="utf-8-sig"))
        found: set[str] = set()
        for node in ast.walk(tree):
            if not isinstance(node, ast.If):
                continue
            for sub in ast.walk(node):
                if not isinstance(sub, ast.Call):
                    continue
                if not isinstance(sub.func, ast.Attribute) or sub.func.attr != "append":
                    continue
                if not sub.args or not isinstance(sub.args[0], ast.Dict):
                    continue
                title = _literal(sub.args[0], "title")
                if title:
                    found.add(title)
        return found

    def test_the_six_items_are_never_conditionally_appended(self) -> None:
        offenders = self._conditional_item_titles() & self.MUST_ALWAYS_RENDER
        self.assertEqual(offenders, set(), f"这些检查项又变回「空数据即消失」: {sorted(offenders)}")

    def test_empty_text_differs_by_collection_status(self) -> None:
        """采集成功与未采集的空数据文案必须不同 —— 不能把「没查」说成「查了没有」。"""
        from plugins.sqlserver.presentation import _blank_note

        confirmed = "已检查，未发现。"
        self.assertEqual(_blank_note({"status": "success_no_data"}, confirmed), confirmed)
        for status in ("not_evaluated", "skipped", "permission_denied", "error", None):
            got = _blank_note({"status": status}, confirmed)
            with self.subTest(status=status):
                self.assertNotEqual(got, confirmed)
                self.assertIn("未完成评价", got)
                self.assertNotIn("已检查", got)

    def test_the_six_items_carry_their_own_empty_text(self) -> None:
        """这 6 项必须自带 empty_text，否则空数据时会退回通用措辞、说不清是哪种空。"""
        source = PRESENTATION.read_text(encoding="utf-8-sig")
        for title in sorted(self.MUST_ALWAYS_RENDER):
            with self.subTest(item=title):
                self.assertIn(f'"title": "{title}"', source)
        self.assertGreaterEqual(source.count("_blank_note("), len(self.MUST_ALWAYS_RENDER))


class MultiRuleEvidenceLabelTests(unittest.TestCase):
    """一个检查项命中多条规则时，每条依据与建议必须自带主语。

    实测反例：「max server memory=…，物理内存=4095 MB；当前值=5；」里那个「当前值=5」
    是 cost threshold 的，不加主语读者会读成最大内存配了 5 MB。
    """

    @staticmethod
    def _run(findings: list[dict]) -> dict:
        from plugins.sqlserver.presentation import _inject_analysis

        sections = [{"section_id": "config", "title": "实例关键配置", "items": [{"title": "关键配置项"}]}]
        _inject_analysis(sections, findings)
        return sections[0]["items"][0]["analysis"]

    def test_each_finding_keeps_its_subject(self) -> None:
        analysis = self._run([
            {"rule_id": "INSTANCE.MAX_MEMORY", "title": "最大内存配置过高",
             "object_name": "max server memory", "severity": "high",
             "evidence": "max server memory=2147483647 MB", "recommendation": "调整最大内存。"},
            {"rule_id": "INSTANCE.COST_THRESHOLD", "title": "并行开销阈值偏低",
             "object_name": "cost threshold for parallelism", "severity": "medium",
             "evidence": "当前值=5", "recommendation": "提高并行开销阈值。"},
        ])
        self.assertIn("cost threshold for parallelism：当前值=5", "；".join(analysis["evidence"]))
        self.assertIn("cost threshold for parallelism：提高并行开销阈值。", analysis["recommendation"])

    def test_single_finding_is_not_labeled(self) -> None:
        """只命中一条规则时不加主语前缀 —— 加了反而是噪音。

        输入与上面那条测试逐字段相同（同规则、同 object_name、同 evidence），
        只差条数，两条合起来才是干净的控制变量：条数一个不影响措辞，一个必须加主语。
        """
        analysis = self._run([
            {"rule_id": "INSTANCE.COST_THRESHOLD", "title": "并行开销阈值偏低",
             "object_name": "cost threshold for parallelism", "severity": "medium",
             "evidence": "当前值=5", "recommendation": "提高并行开销阈值。"},
        ])
        self.assertEqual(analysis["evidence"], ["当前值=5"])
        self.assertEqual(analysis["recommendation"], "提高并行开销阈值。")


if __name__ == "__main__":
    unittest.main()
